"""Seed SQLite from the provided Excel workbook. Idempotent: rebuilds from scratch."""
import os
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from perennials.db import connect, init_db  # noqa: E402

DEFAULT_XLSX = os.path.join(os.path.dirname(__file__), "..", "data", "perennials_sample_data.xlsx")


def seed(xlsx_path=None, db_path=None):
    xlsx_path = xlsx_path or DEFAULT_XLSX
    conn = connect(db_path)
    # full rebuild so re-running is safe
    tables = ["erp_items", "erp_lots", "pricing_rules", "crm_accounts", "crm_contacts",
              "contact_emails", "raw_inbound", "extractions", "match_decisions",
              "crm_opportunities", "erp_sample_orders", "erp_sample_order_lines",
              "crm_quotes", "crm_quote_lines", "erp_reservations", "review_queue", "audit_log"]
    for t in tables:
        conn.execute("DROP TABLE IF EXISTS %s" % t)
    init_db(conn)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    ws = wb["Fabric Catalog"]
    for row in list(ws.iter_rows(values_only=True))[1:]:
        sku, fabric, color, status, price, lot, on_hand, reserved = row
        conn.execute(
            "INSERT OR IGNORE INTO erp_items VALUES (?,?,?,?,?)", (sku, fabric, color, status, price))
        conn.execute(
            "INSERT INTO erp_lots(sku,dye_lot,yards_on_hand,yards_reserved) VALUES (?,?,?,?)",
            (sku, lot, on_hand, reserved))

    ws = wb["Pricing Rules"]
    for row in list(ws.iter_rows(values_only=True))[1:]:
        conn.execute("INSERT OR REPLACE INTO pricing_rules VALUES (?,?,?)", row)

    ws = wb["Existing CRM Records"]
    for row in list(ws.iter_rows(values_only=True))[1:]:
        rec_type, rid, account_id, name, email, typ, status = row
        if rec_type == "Account":
            conn.execute("INSERT OR REPLACE INTO crm_accounts(id,name,type,status) VALUES (?,?,?,?)",
                         (rid, name, typ, status))
        elif rec_type == "Contact":
            conn.execute("INSERT OR REPLACE INTO crm_contacts(id,account_id,name,email) VALUES (?,?,?,?)",
                         (rid, account_id, name, email))
            if email:
                conn.execute("INSERT OR IGNORE INTO contact_emails VALUES (?,?,1)", (rid, email.lower()))

    ws = wb["Inbound Requests"]
    for row in list(ws.iter_rows(values_only=True))[1:]:
        conn.execute("INSERT OR REPLACE INTO raw_inbound(id,type,channel,received,name,company,email,message) "
                     "VALUES (?,?,?,?,?,?,?,?)", row)
    conn.commit()
    counts = {t: conn.execute("SELECT COUNT(*) c FROM %s" % t).fetchone()["c"]
              for t in ["erp_items", "erp_lots", "crm_accounts", "crm_contacts", "raw_inbound"]}
    conn.close()
    return counts


if __name__ == "__main__":
    print("Seeded:", seed(sys.argv[1] if len(sys.argv) > 1 else None))
